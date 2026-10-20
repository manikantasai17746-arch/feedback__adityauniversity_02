"""Import the timetable workbook into the database.

Builds, from the real source data:

* ``assignments``  - unchanged legacy shape, plus timetable/period/semester/year
* ``branches``, ``semesters``, ``academic_years``, ``timetables``
* ``users`` (teachers + students), ``class_sessions``, ``enrollments``

Everything is idempotent: re-running only inserts what is missing.
"""
import os
import re
import time
from collections import defaultdict

import openpyxl
from werkzeug.security import generate_password_hash

from . import db


def _row_id(row):
    """Return the primary-key value from a sqlite3.Row or mapping."""
    if row is None:
        return None
    try:
        return row["id"]
    except Exception:
        return row[0]


HEADER_ROWS = 4  # rows 1..4 are titles/headers in the source workbook


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _branch_name(title, row3):
    """Best-effort human readable branch name from the sheet title row."""
    if row3:
        text = str(row3)
        if ',' in text:
            tail = text.split(',', 1)[1].strip()
            if tail:
                return tail
    return title


def ensure_reference_data(conn, semester_number, academic_year_name):
    conn.execute('INSERT OR IGNORE INTO academic_years(name) VALUES(?)', (academic_year_name,))
    for n in range(1, 9):
        conn.execute('INSERT OR IGNORE INTO semesters(number, label) VALUES(?,?)',
                     (n, f'Semester {n}'))
    ay = conn.execute('SELECT id FROM academic_years WHERE name=?',
                      (academic_year_name,)).fetchone()[0]
    sem = conn.execute('SELECT id FROM semesters WHERE number=?',
                       (semester_number,)).fetchone()[0]
    return ay, sem


def ensure_admin(conn, email, password):
    """Create the first admin only when none exists. Never overwrites an existing admin.

    Returns the admin user id, or None when credentials are missing and no admin
    has been created yet (caller decides whether to raise).
    """
    row = conn.execute("SELECT id FROM users WHERE role='admin' LIMIT 1").fetchone()
    if row:
        return _row_id(row)
    if not email or not password:
        return None
    email = email.strip().lower()
    conn.execute(
        'INSERT INTO users(role, username, email, name, password_hash, is_active, '
        'must_change_password) VALUES(?, ?, ?, ?, ?, 1, 0)',
        ('admin', email, email, 'Administrator', generate_password_hash(password)))
    row = conn.execute("SELECT id FROM users WHERE role='admin' LIMIT 1").fetchone()
    return _row_id(row)


def set_admin_credentials(conn, email, password):
    """Explicitly create or replace the first admin's login credentials.

    Safe to re-run: updates username, email, and password_hash for the first
    admin row, or creates one when the table has no admin yet.
    """
    if not email or not password:
        raise RuntimeError(
            'Set ADMIN_EMAIL and ADMIN_PASSWORD in the environment before syncing admin credentials.')
    email = email.strip().lower()
    row = conn.execute(
        "SELECT id FROM users WHERE role='admin' ORDER BY id LIMIT 1").fetchone()
    if row:
        admin_id = _row_id(row)
        conn.execute(
            'UPDATE users SET username=?, email=?, password_hash=?, is_active=1, '
            'must_change_password=0 WHERE id=?',
            (email, email, generate_password_hash(password), admin_id))
    else:
        admin_id = ensure_admin(conn, email, password)
        if admin_id is None:
            raise RuntimeError('Failed to create admin account.')
    conn.commit()
    return admin_id


def reset_student_teacher_passwords(conn, teacher_password, student_password):
    """Reset all student and teacher accounts to their configured defaults."""
    if not teacher_password or not student_password:
        raise RuntimeError(
            'Set DEFAULT_TEACHER_PASSWORD and DEFAULT_STUDENT_PASSWORD before resetting accounts.')
    teacher_hash = generate_password_hash(teacher_password)
    student_hash = generate_password_hash(student_password)
    try:
        teacher_count = conn.execute(
            "SELECT COUNT(*) FROM users WHERE role='teacher'").fetchone()[0]
        student_count = conn.execute(
            "SELECT COUNT(*) FROM users WHERE role='student'").fetchone()[0]
        conn.execute(
            "UPDATE users SET password_hash=?, must_change_password=1 WHERE role='teacher'",
            (teacher_hash,))
        conn.execute(
            "UPDATE users SET password_hash=? WHERE role='student'",
            (student_hash,))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return teacher_count, student_count


# --------------------------------------------------------------------------- #
# Reading + validating a timetable workbook
# --------------------------------------------------------------------------- #
REQUIRED_HEADERS = {
    'timetable': ('time table', 'timetable'),
    'course_code': ('course code',),
    'course_name': ('course name',),
    'emp_code': ('emp.code', 'emp code', 'empcode'),
    'emp_name': ('emp.name', 'emp name', 'empname'),
    'roll': ('roll.no', 'roll no', 'rollno', 'roll'),
}
ACADEMIC_YEAR_RE = re.compile(r'^\d{4}-\d{2}$')


class ImportErrorReport(ValueError):
    """Raised when a workbook cannot be imported; ``problems`` lists why."""

    def __init__(self, problems):
        self.problems = list(problems)
        super().__init__('; '.join(self.problems[:5]))


def claim_timetable_import(conn, fingerprint, semester, academic_year, filename,
                           now=None, stale_after=3600):
    """Atomically claim an import, returning (state, import_id)."""
    now = int(time.time() if now is None else now)
    row = conn.execute(
        """INSERT INTO timetable_imports(
               fingerprint, semester, academic_year, filename, status, started_ts)
           VALUES(?,?,?,?, 'importing', ?)
           ON CONFLICT(fingerprint, semester, academic_year) DO NOTHING
           RETURNING id""",
        (fingerprint, semester, academic_year, filename, now)).fetchone()
    if row:
        conn.commit()
        return 'claimed', row[0]

    existing = conn.execute(
        """SELECT id, status, started_ts FROM timetable_imports
           WHERE fingerprint=? AND semester=? AND academic_year=?""",
        (fingerprint, semester, academic_year)).fetchone()
    if not existing:
        raise RuntimeError('Could not determine the status of the workbook import.')
    if existing[1] == 'completed':
        return 'completed', None
    if existing[2] < now - stale_after:
        row = conn.execute(
            """UPDATE timetable_imports SET filename=?, started_ts=?
               WHERE id=? AND status='importing' AND started_ts < ?
               RETURNING id""",
            (filename, now, existing[0], now - stale_after)).fetchone()
        if row:
            conn.commit()
            return 'claimed', row[0]
    return 'importing', None


def complete_timetable_import(conn, import_id):
    conn.execute("UPDATE timetable_imports SET status='completed' WHERE id=?",
                 (import_id,))
    conn.commit()


def release_timetable_import(conn, import_id):
    conn.execute("DELETE FROM timetable_imports WHERE id=? AND status='importing'",
                 (import_id,))
    conn.commit()


def _norm_header(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().lower()


def _clean(value):
    """Cell -> trimmed string; whole-number floats lose the trailing '.0'."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _find_header(rows):
    """Return (index_in_rows, {field: column_index}) for the header row."""
    for i, row in enumerate(rows):
        cells = [_norm_header(c) for c in row]
        cols = {}
        for field, names in REQUIRED_HEADERS.items():
            for j, c in enumerate(cells):
                if c in names:
                    cols[field] = j
                    break
        if len(cols) == len(REQUIRED_HEADERS):
            return i, cols
    return None, {}


def read_workbook(source):
    """Parse a timetable workbook (path or file-like object).

    Returns ``(raw, branch_names, report)`` where ``raw`` is a list of
    ``(dept, timetable, course_code, course_name, emp_code, emp_name, roll)``
    tuples.  ``report`` carries ``sheets``, ``skipped_rows`` and ``warnings``.
    Raises :class:`ImportErrorReport` if nothing usable is found.
    """
    try:
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportErrorReport(
            [f'This is not a readable .xlsx workbook ({type(exc).__name__}).'])

    raw, branch_names = [], {}
    report = {'sheets': [], 'skipped_rows': 0, 'warnings': []}
    try:
        for ws in wb:
            head = list(ws.iter_rows(min_row=1, max_row=12, values_only=True))
            hidx, cols = _find_header(head)
            if hidx is None:
                report['warnings'].append(
                    f'Sheet "{ws.title}" skipped: no header row with '
                    f'Time Table / Course Code / Course Name / Emp.Code / '
                    f'Emp.Name / Roll.No.')
                continue
            title_row = head[2][0] if len(head) >= 3 and head[2] and hidx >= 3 else None
            branch_names[ws.title] = _branch_name(ws.title, title_row)
            count = 0
            for r in ws.iter_rows(min_row=hidx + 2, values_only=True):
                if not r or all(c is None or _clean(c) == '' for c in r):
                    continue
                vals = {f: _clean(r[j]) if j < len(r) else '' for f, j in cols.items()}
                if not all(vals[f] for f in ('roll', 'emp_code', 'course_code', 'timetable')):
                    report['skipped_rows'] += 1
                    continue
                raw.append((ws.title, vals['timetable'], vals['course_code'],
                            vals['course_name'] or vals['course_code'],
                            vals['emp_code'], vals['emp_name'] or vals['emp_code'],
                            vals['roll'].upper()))
                count += 1
            report['sheets'].append((ws.title, count))
    finally:
        wb.close()

    if not raw:
        raise ImportErrorReport(
            report['warnings'] or ['No timetable rows were found in this workbook.'])
    if report['skipped_rows']:
        report['warnings'].append(
            f"{report['skipped_rows']} row(s) skipped because Roll.No, Emp.Code, "
            f"Course Code or Time Table was empty.")
    return raw, branch_names, report


def summarize(raw, report):
    """Counts shown on the 'Check file' preview (nothing is written)."""
    return {
        'rows': len(raw),
        'branches': len({r[0] for r in raw}),
        'timetables': len({(r[0], r[1]) for r in raw}),
        'subjects': len({r[2] for r in raw}),
        'teachers': len({r[4] for r in raw}),
        'students': len({r[6] for r in raw}),
        'sheets': report['sheets'],
        'warnings': report['warnings'],
    }


def validate_semester_year(semester_number, academic_year):
    try:
        semester_number = int(semester_number)
    except (TypeError, ValueError):
        raise ImportErrorReport(['Semester must be a number from 1 to 8.'])
    if not 1 <= semester_number <= 8:
        raise ImportErrorReport(['Semester must be a number from 1 to 8.'])
    academic_year = (academic_year or '').strip()
    if not ACADEMIC_YEAR_RE.match(academic_year):
        raise ImportErrorReport(['Academic year must look like 2025-26.'])
    return semester_number, academic_year


# --------------------------------------------------------------------------- #
# Main import
# --------------------------------------------------------------------------- #
def import_timetable(app, semester_number=3, academic_year=None, progress=print,
                     source=None, raw_data=None):
    """Import a timetable workbook for one semester / academic year.

    ``source`` is a path or file-like object (defaults to the configured
    workbook).  Callers that already parsed the file may pass ``raw_data`` as
    the ``(raw, branch_names)`` pair to avoid reading it twice.
    """
    academic_year = academic_year or os.environ.get('ACADEMIC_YEAR') or '2025-26'
    semester_number, academic_year = validate_semester_year(semester_number, academic_year)
    conn = db.get_db()

    # --- 1. read workbook -------------------------------------------------- #
    if raw_data is not None:
        raw, branch_names = raw_data
        label = 'uploaded workbook'
    else:
        source = source or app.config['XLSX_PATH']
        if isinstance(source, str) and not os.path.exists(source):
            raise FileNotFoundError(f'Timetable workbook not found: {source}')
        raw, branch_names, _report = read_workbook(source)
        label = os.path.basename(source) if isinstance(source, str) else 'uploaded workbook'
    progress(f'Read {len(raw)} timetable rows from {label}')

    # --- 2. period index per student (position within their own timetable) -- #
    subject_period = {}                       # (dept, roll, cc) -> 1-based period
    seen_order = defaultdict(dict)            # (dept, roll) -> {cc: period}
    for dept, tt, cc, cn, emp, empname, roll in raw:
        bucket = seen_order[(dept, roll)]
        if cc not in bucket:
            bucket[cc] = len(bucket) + 1
    for (dept, roll), bucket in seen_order.items():
        for cc, p in bucket.items():
            subject_period[(dept, roll, cc)] = p

    # --- 3. reference rows ------------------------------------------------- #
    ay_id, sem_id = ensure_reference_data(conn, semester_number, academic_year)
    progress(f'Academic year: {academic_year} (id={ay_id}), semester {semester_number}')

    # 3a. branches
    branch_id = {}
    for dept, name in branch_names.items():
        conn.execute('INSERT OR IGNORE INTO branches(code, name) VALUES(?,?)', (dept, name))
        branch_id[dept] = conn.execute('SELECT id FROM branches WHERE code=?',
                                       (dept,)).fetchone()[0]

    # 3b. legacy assignments (original shape + new columns)
    progress('Importing assignments...')
    conn.executemany(
        """INSERT OR IGNORE INTO assignments(
               roll, dept, course_code, course_name, emp_code, emp_name,
               timetable, period_no, semester, academic_year)
           VALUES(?,?,?,?,?,?,?,?,?,?)""",
        [(roll, dept, cc, cn, emp, empname, tt,
          subject_period.get((dept, roll, cc)), semester_number, academic_year)
         for (dept, tt, cc, cn, emp, empname, roll) in raw])

    # 3c. teacher users
    teacher_hash = generate_password_hash(app.config['DEFAULT_TEACHER_PASSWORD'])
    teacher_users = {}
    for dept, tt, cc, cn, emp, empname, roll in raw:
        teacher_users.setdefault(emp, empname)
    conn.executemany(
        """INSERT OR IGNORE INTO users(role, username, name, emp_code, password_hash)
           VALUES('teacher', ?, ?, ?, ?)""",
        [(emp, name, emp, teacher_hash) for emp, name in teacher_users.items()])
    teacher_id = {row[0]: row[1] for row in
                  conn.execute("SELECT emp_code, id FROM users WHERE role='teacher'")}
    progress(f'Teachers: {len(teacher_id)}')

    # 3d. student users
    student_hash = generate_password_hash(app.config['DEFAULT_STUDENT_PASSWORD'])
    students = sorted({roll for (*_, roll) in raw})
    conn.executemany(
        """INSERT OR IGNORE INTO users(role, username, name, roll, password_hash)
           VALUES('student', ?, ?, ?, ?)""",
        [(roll, roll, roll, student_hash) for roll in students])
    student_id = {row[0]: row[1] for row in
                  conn.execute("SELECT roll, id FROM users WHERE role='student'")}
    progress(f'Students: {len(student_id)}')

    # 3e. timetables
    tt_pairs = sorted({(dept, tt) for dept, tt, *_ in raw})
    timetable_id = {}
    for dept, tt in tt_pairs:
        conn.execute(
            """INSERT OR IGNORE INTO timetables(
                   branch_id, semester_id, academic_year_id, code, name)
               VALUES(?,?,?,?,?)""", (branch_id[dept], sem_id, ay_id, tt, tt))
        timetable_id[(dept, tt)] = conn.execute(
            'SELECT id FROM timetables WHERE branch_id=? AND semester_id=? AND '
            'academic_year_id=? AND code=?', (branch_id[dept], sem_id, ay_id, tt)).fetchone()[0]
    progress(f'Timetables: {len(timetable_id)}')

    # 3f. class sessions (timetable + subject + teacher) with a stable period
    sess_period, sess_names = {}, {}
    for dept, tt, cc, cn, emp, empname, roll in raw:
        k = (dept, tt, cc, emp)
        p = subject_period.get((dept, roll, cc))
        if p is not None and (k not in sess_period or p < sess_period[k]):
            sess_period[k] = p
        sess_names[k] = cn
    conn.executemany(
        """INSERT OR IGNORE INTO class_sessions(
               timetable_id, subject_code, subject_name, teacher_user_id, period_no)
           VALUES(?,?,?,?,?)""",
        [(timetable_id[(dept, tt)], cc, sess_names[(dept, tt, cc, emp)],
          teacher_id.get(emp), period)
         for (dept, tt, cc, emp), period in sess_period.items()])
    session_id = {}
    for row in conn.execute(
            """SELECT cs.id, t.branch_id, t.code AS tt, cs.subject_code, u.emp_code
               FROM class_sessions cs
               JOIN timetables t ON t.id = cs.timetable_id
               LEFT JOIN users u ON u.id = cs.teacher_user_id"""):
        dept = next((d for d, n in branch_id.items() if n == row['branch_id']), None)
        session_id[(dept, row['tt'], row['subject_code'], row['emp_code'])] = row['id']
    progress(f'Class sessions: {len(session_id)}')

    # 3g. enrollments (each student belongs to exactly one timetable)
    student_tt = {}
    for dept, tt, cc, cn, emp, empname, roll in raw:
        student_tt[roll] = (dept, tt)
    conn.executemany(
        """INSERT OR IGNORE INTO enrollments(
               student_user_id, timetable_id, semester_id, academic_year_id)
           VALUES(?,?,?,?)""",
        [(student_id[roll], timetable_id[key], sem_id, ay_id)
         for roll, key in student_tt.items()])
    progress(f'Enrollments: {len(student_tt)}')

    # 3h. eligible student population per teaching session.
    # A student is eligible for a session only when the timetable assigns that
    # exact (timetable, subject, teacher) triple to them.
    conn.executemany(
        'INSERT OR IGNORE INTO session_students(session_id, student_user_id) VALUES(?,?)',
        [(session_id[(dept, tt, cc, emp)], student_id[roll])
         for (dept, tt, cc, cn, emp, empname, roll) in raw
         if (dept, tt, cc, emp) in session_id])
    n_elig = conn.execute('SELECT COUNT(*) FROM session_students').fetchone()[0]
    progress(f'Session eligibility rows: {n_elig}')

    conn.commit()
    return {
        'rows': len(raw), 'branches': len(branch_id), 'timetables': len(timetable_id),
        'sessions': len(session_id), 'teachers': len(teacher_id),
        'students': len(student_id), 'enrollments': len(student_tt),
    }


def bootstrap_if_empty(app):
    """Create the admin account (if missing) and import the workbook on a fresh DB.

    Admin is only created when ADMIN_EMAIL and ADMIN_PASSWORD are set and no
    admin row exists yet. Existing admin passwords are never changed here —
    use ``scripts/sync_admin_credentials.py`` to reset them deliberately.
    """
    conn = db.get_db()
    admin_id = ensure_admin(conn, app.config.get('ADMIN_EMAIL'),
                            app.config.get('ADMIN_PASSWORD'))
    conn.commit()
    if admin_id is None:
        # No credentials and no admin: surface a clear operational message.
        # Do not crash the process — login will simply fail until credentials
        # are configured and the sync script (or a restart with env set) runs.
        app.logger.error(
            'No admin account exists and ADMIN_EMAIL/ADMIN_PASSWORD are not set. '
            'Set both environment variables and run: python scripts/sync_admin_credentials.py')
    count_row = conn.execute('SELECT COUNT(*) FROM assignments').fetchone()
    count = 0
    if count_row is not None:
        try:
            count = int(count_row[0])
        except Exception:
            count = int(list(count_row)[0])
    if count == 0 and os.path.exists(app.config['XLSX_PATH']):
        try:
            import_timetable(app)
        except Exception as exc:  # pragma: no cover - surface but don't crash boot
            app.logger.exception('Timetable import failed: %s', exc)
